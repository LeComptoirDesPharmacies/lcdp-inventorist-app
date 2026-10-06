import unittest

from openpyxl import load_workbook

from business.mappers.excel_lines_mapper import LaboratoryExcelLinesMapper, ProductExcelLinesMapper, \
    DrugstoreExcelLinesMapper
from tests.constant import LABORATORY_SALE_OFFER_EXCEL, IMPORT_PRODUCT_EXCEL, DRUGSTORE_SALE_OFFER_EXCEL


class TestExcelMapperHeaders(unittest.TestCase):
    """A column whose header no longer matches the template is silently ignored by the mapper:
    every column without a default value must be found in the fixture workbook (a copy of the
    template headers)."""

    def assert_columns_are_in_workbook(self, mapper):
        parameters = mapper.get_parameters()
        wb = load_workbook(mapper.excel_path, read_only=True)
        try:
            header_row = next(wb[parameters.sheet_name].iter_rows(
                min_row=int(parameters.header_line), max_row=int(parameters.header_line), values_only=True))
        finally:
            wb.close()
        headers = [header for header in header_row if header is not None]

        missing = [column.excel_column_name for column in mapper.excel_mapper
                   if column.default_value is None and column.excel_column_name not in headers]
        self.assertEqual([], missing)

    def test_laboratory_sale_offer_columns(self):
        self.assert_columns_are_in_workbook(LaboratoryExcelLinesMapper(LABORATORY_SALE_OFFER_EXCEL))

    def test_product_columns(self):
        self.assert_columns_are_in_workbook(ProductExcelLinesMapper(IMPORT_PRODUCT_EXCEL))

    def test_drugstore_sale_offer_columns(self):
        self.assert_columns_are_in_workbook(DrugstoreExcelLinesMapper(DRUGSTORE_SALE_OFFER_EXCEL))
